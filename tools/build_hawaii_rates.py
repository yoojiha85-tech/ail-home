#!/usr/bin/env python3
"""하와이 호텔 요금표(엑셀, 차수별) -> data/hawaii-rates-<차수>.js 변환 도구.

사용법: python3 -I tools/build_hawaii_rates.py <엑셀경로> <차수라벨> <요금표기준일>
예) python3 -I tools/build_hawaii_rates.py 요금표.xlsx 21 2026-10-07
조건이 애매한 요금은 자동으로 'review'(확인 필요) 표시를 달고, 날짜가 이상한 행은 'error'로 분리합니다.
"""
import json, re, sys, datetime
import openpyxl

SKIP = ('★목차', '25년 오아후')
REVIEW_KW = ['black', '블랙', '미니멈', '최소', '박 이상', '박 예약', '신규', '콤보', 'upgrade', '업그레이드',
             '무료', 'stay ', 'pay ', '특전', '개런티', '별도문의', '조건']

def clean(v):
    return re.sub(r'\s+', ' ', str(v)).strip() if v is not None else ''

def parse_dates(s):
    # 연도가 빠진 날짜(예: "12/20 ~ 12/31/2026")는 바로 뒤에 나오는 연도를 따라갑니다.
    toks = re.findall(r'(\d{1,2})/(\d{1,2})(?:/(\d{2,4}))?', s or '')
    out = []
    for i, (m, d, y) in enumerate(toks):
        if not y:
            y = next((t[2] for t in toks[i + 1:] if t[2]), '')
        if not y:
            out.append(None); continue
        y = int(y); y = y + 2000 if y < 100 else y
        try: out.append(datetime.date(y, int(m), int(d)))
        except ValueError: out.append(None)
    return out

def periods(s):
    ds = parse_dates(s)
    res = []
    for i in range(0, len(ds) - 1, 2):
        res.append((ds[i], ds[i + 1]))
    return res, len(ds)

def build(xlsx, label, asof):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    hotels = []
    for ws in wb.worksheets:
        if ws.title.startswith(SKIP):
            continue
        name_raw = clean(ws['A2'].value)
        if not name_raw:
            continue
        island = '마우이' if ws.title.startswith('마우이') else ('빅아일랜드' if ws.title.startswith('빅아일랜드') else '오아후')
        grade = re.search(r'★\s*([\d.]+)', name_raw) or re.search(r'\(([\d.]+)성\)', ws.title)
        h = {'sheet': ws.title.strip(), 'name': re.sub(r'\s*★.*$', '', name_raw).strip(),
             'island': island, 'grade': grade.group(1) if grade else '', 'notes': [], 'cancel': [], 'blocks': []}
        cols = None; bf = None; header_note = []
        pending = []; last_was_rate = False; block = None; period = None; in_cancel = False
        for row in ws.iter_rows(min_row=3, max_col=12):
            cells = {c.column - 1: c.value for c in row if c.value is not None}
            a = clean(cells.get(0))
            texts = [clean(v) for v in cells.values() if isinstance(v, str) and re.search(r'\w', clean(v))]
            if '예약 및 취소규정' in a or '예약및취소규정' in a.replace(' ', ''):
                in_cancel = True
                continue
            if in_cancel:
                if texts: h['cancel'].append(' / '.join(texts)[:400])
                continue
            # header row
            hdr = [i for i, v in cells.items() if clean(v) in ('기간', '적용 기간')]
            if hdr:
                pc = hdr[0]
                fc = next((i for i, v in cells.items() if '리조트피' in clean(v)), None)
                rc = next((i for i, v in cells.items() if '1박' in clean(v)), None)
                htxt = clean(cells.get(rc)).replace(' ', '') if rc is not None else ''
                bf = False if '불포함' in htxt else ('조식' in htxt)
                cols = (pc, fc, rc)
                extra = [clean(v) for i, v in cells.items() if i > (rc or 0) and isinstance(v, str) and clean(v)]
                if last_was_rate:
                    # 같은 프로그램의 조식 포함/불포함 표처럼 제목 없이 표가 이어지면 앞 표의 제목·조건을 이어받음
                    pending = [] if a else list(h['blocks'][-1]['notes'])
                    block = None
                if a: pending.append(a)
                pending += extra
                last_was_rate = False; period = None; block = None
                continue
            rate = cells.get(cols[2]) if cols else None
            if cols and isinstance(rate, (int, float)) and a:
                if block is None:
                    block = {'title': (pending[0] if pending else '').split('[')[0].strip() or '요금',
                             'notes': list(pending), 'bf': bf, 'rows': []}
                    h['blocks'].append(block)
                p = clean(cells.get(cols[0])) if cols[0] in cells else ''
                if p:
                    period = p
                fee = clean(cells.get(cols[1])) if cols[1] is not None else ''
                block['rows'].append({'room': a, 'period': period or '', 'rate': rate, 'fee_incl': bool(fee)})
                last_was_rate = True
                continue
            if texts:
                if cols and cols[0] in cells and re.search(r'\d{1,2}/\d{1,2}', clean(cells.get(cols[0]))):
                    period = clean(cells.get(cols[0]))   # 시즌명 줄에 기간이 같이 적힌 경우(예: Season B)
                if last_was_rate:
                    pending = []; block = None
                    last_was_rate = False
                if cols is None:
                    pending += texts
                    h['notes'] += [t for t in texts if t.startswith('※') or '포터리지' in t]
                else:
                    pending += texts
        # 호텔 공통 안내 (A3 등 첫 줄 메모)
        h['notes'] = [n for n in h['notes']]
        hotels.append(h)

    out_hotels = []
    stats = {'rows': 0, 'ok': 0, 'review': 0, 'error': 0}
    for h in hotels:
        top = ' '.join(h['notes'])
        portage = re.search(r'포터리지[^$]*\$\s*(\d+)', top)
        oh = {'id': h['sheet'], 'name': h['name'], 'island': h['island'], 'grade': h['grade'],
              'portage_usd': int(portage.group(1)) if portage else 0,
              'notes': h['notes'], 'cancel': h['cancel'], 'programs': []}
        for b in h['blocks']:
            text = ' '.join(b['notes'])
            low = text.lower()
            win = re.search(r'booking\s*window\s*:?\s*([\d/ ]*)~\s*([\d/ ]+)', text, re.I)
            wstart = wend = None
            if win:
                s1 = parse_dates(win.group(1)); s2 = parse_dates(win.group(2))
                wstart = s1[0] if s1 else None; wend = s2[0] if s2 else None
            minn = None
            m = re.search(r'미니멈\s*(\d)\s*박|(\d)\s*박\s*미니멈|최소\s*(\d)\s*박|(\d)\s*박\s*이상\s*예약', text)
            if m: minn = int(next(g for g in m.groups() if g))
            bo = [clean(n) for n in b['notes'] if re.search(r'black|블랙', n, re.I)]
            review = any(k in low for k in REVIEW_KW)
            prog = {'title': b['title'], 'breakfast': b['bf'], 'notes': b['notes'],
                    'window': [wstart.isoformat() if wstart else None, wend.isoformat() if wend else None],
                    'min_nights': minn, 'blackout': bo, 'review': review, 'rows': []}
            for r in b['rows']:
                ps, nd = periods(r['period'])
                issues = []
                if not ps: issues.append('기간을 읽을 수 없음')
                elif nd % 2: issues.append('기간 날짜 개수가 홀수')
                for s, e in ps:
                    if s is None or e is None: issues.append('존재하지 않는 날짜')
                    elif s > e: issues.append('시작일이 종료일보다 늦음')
                    elif s.year < 2026 or (e - s).days > 400: issues.append('연도/기간 이상(%s~%s)' % (s, e))
                stats['rows'] += 1
                row_review = review or any(k in r['room'].lower() for k in REVIEW_KW)
                rm = re.search(r'(\d)\s*박\s*미니멈|미니멈\s*(\d)\s*박', r['period'])
                rr = {'room': r['room'], 'rate': r['rate'], 'fee_incl': r['fee_incl'], 'raw': r['period'],
                      'min_nights': int(next(g for g in rm.groups() if g)) if rm else None,
                      'periods': [[s.isoformat(), e.isoformat()] for s, e in ps if s and e and not issues]}
                if row_review: rr['review'] = True
                if issues:
                    rr['error'] = '; '.join(sorted(set(issues))); stats['error'] += 1
                else:
                    stats['review' if row_review else 'ok'] += 1
                prog['rows'].append(rr)
            oh['programs'].append(prog)
        out_hotels.append(oh)
    return {'label': label, 'asof': asof, 'generated': datetime.date.today().isoformat(), 'hotels': out_hotels}, stats

if __name__ == '__main__':
    xlsx, label, asof = sys.argv[1:4]
    data, stats = build(xlsx, label, asof)
    path = 'data/hawaii-rates-%s.js' % label
    with open(path, 'w', encoding='utf-8') as f:
        f.write('/* 자동 생성 파일 - tools/build_hawaii_rates.py 로 만든 요금표 데이터입니다. 직접 수정하지 마세요. */\n')
        f.write('window.HAWAII_RATES = ' + json.dumps(data, ensure_ascii=False, separators=(',', ':')) + ';\n')
    print(path, stats, 'hotels', len(data['hotels']))
